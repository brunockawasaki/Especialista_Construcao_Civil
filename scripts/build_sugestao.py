# -*- coding: utf-8 -*-
import sys, re, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from aplicar import pontuar
from score_base import ADJ, W, PROC_NOTES, rnd, FACTORS

OUT = sys.argv[1]
CUT_NORMAL, CUT_ALTA, CUT_URG = 2.5, 3.3, 3.9

rows, byid, man = pontuar()

def situacao(st):
    return 'Concluída' if st == 'concluído' else 'Cancelada' if st == 'cancelado' else 'Aberta'

# ---------- calculo em Python (para conferir com o Excel) ----------
for o in rows:
    s = o['s']; n = o['nota']
    o['catn'] = 4 if n >= CUT_URG else 3 if n >= CUT_ALTA else 2 if n >= CUT_NORMAL else 1
    veto = 0
    if s['R'] == 5: veto = max(veto, 3)
    if s['cef'] == 'Sim': veto = max(veto, 4)
    o['veto'] = veto
    o['cat'] = max(o['catn'], veto)
    o['prio'] = ['Baixa', 'Normal', 'Alta', 'Urgente'][o['cat'] - 1]
    o['sit'] = situacao(o['t']['status'])
    o['resp1'] = o['t']['assignees'][0] if o['t']['assignees'] else '(sem responsável)'
    o['resps'] = ', '.join(o['t']['assignees']) or '(sem responsável)'
    o['key0'] = o['cat'] * 1000 + o['nota'] * 10 - s['E'] / 10

# ordem estatica das linhas = ordem geral (chave desc; empate = ordem original)
rows.sort(key=lambda o: -o['key0'])
for i, o in enumerate(rows, 1): o['ordem'] = i

LBL = [('R', 'risco de fraude/erro'), ('F', 'impacto financeiro'), ('O', 'obra e entrega'), ('L', 'exigência legal/CEF'), ('D', 'dependência')]
def porque(o):
    s = o['s']
    top = sorted([(s[k], n) for k, n in LBL if s[k] >= 4], reverse=True)
    txt = '; '.join('%s %d' % (n, v) for v, n in top) if top else 'nenhum critério com nota 4 ou 5'
    extra = []
    if s['nat'] == 'Controle': extra.append('é controle (+0,3)')
    if s['nat'] in ('Indicador', 'Treinamento'): extra.append('%s (%s)' % (s['nat'].lower(), ('%+.1f' % ADJ[s['nat']]).replace('.', ',')))
    if o['veto'] > o['catn']: extra.append('veto: %s' % ('risco 5 sobe para Alta' if o['veto'] == 3 else 'para obra/repasse CEF sobe para Urgente'))
    return txt + ('. ' + '; '.join(extra) + '.' if extra else '.')

# ---------- estilos ----------
F = 'Arial'
f_n = Font(name=F, size=10); f_b = Font(name=F, size=10, bold=True)
f_t = Font(name=F, size=14, bold=True, color='1F3864'); f_s = Font(name=F, size=10, italic=True, color='595959')
f_h = Font(name=F, size=10, bold=True, color='FFFFFF'); f_in = Font(name=F, size=10, color='0000FF')
f_l = Font(name=F, size=10, color='0563C1', underline='single'); f_g = Font(name=F, size=9, color='7F7F7F')
f_gl = Font(name=F, size=10, color='8FAADC', underline='single')
fill_h = PatternFill('solid', start_color='1F3864'); fill_in = PatternFill('solid', start_color='FFF2CC')
fill_sec = PatternFill('solid', start_color='D9E1F2'); fill_c = PatternFill('solid', start_color='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top'); ctr = Alignment(horizontal='center', vertical='center', wrap_text=True)

def put(ws, ref, v, font=f_n, fill=None, al=wrap, border=True, fmt=None):
    c = ws[ref]; c.value = v; c.font = font
    if fill: c.fill = fill
    c.alignment = al
    if border: c.border = box
    if fmt: c.number_format = fmt
    return c
def head(ws, row, labels, col0=1, h=None):
    for i, l in enumerate(labels): put(ws, '%s%d' % (get_column_letter(col0 + i), row), l, f_h, fill_h, ctr)
    if h: ws.row_dimensions[row].height = h

wb = Workbook()

# =====================================================================
# PARAMETROS
# =====================================================================
wp = wb.create_sheet('Parametros')
wp.sheet_view.showGridLines = False
for k, v in {'A': 38, 'B': 14, 'C': 14, 'D': 16, 'E': 10, 'F': 10, 'G': 60}.items(): wp.column_dimensions[k].width = v
put(wp, 'A1', 'Parâmetros da matriz', f_t, border=False, al=Alignment(vertical='center'))
put(wp, 'A2', 'Células azuis em fundo amarelo podem ser editadas; a aba Todas recalcula. As abas Urgente e Alta são uma fotografia do cálculo na data de geração.', f_s, border=False)
head(wp, 3, ['Critério', 'Peso (%)'])
for i, (k, n) in enumerate([('R', 'Risco de fraude e erro'), ('F', 'Impacto financeiro'), ('O', 'Impacto na obra e na entrega'), ('L', 'Exigência legal, CEF e compliance'), ('D', 'Dependência')]):
    put(wp, 'A%d' % (4 + i), n, f_b, fill_sec); put(wp, 'B%d' % (4 + i), int(W[k] * 100), f_in, fill_in, ctr)
put(wp, 'A9', 'Soma dos pesos', f_b, fill_c); put(wp, 'B9', '=SUM(B4:B8)', f_b, fill_c, ctr)
put(wp, 'C9', '=IF(B9=100,"ok","os pesos devem somar 100")', f_n, fill_c, border=False)
put(wp, 'A11', 'Cortes das categorias (nota mínima)', f_b, border=False)
put(wp, 'A12', '2 · Normal (≥)', f_b, fill_sec); put(wp, 'B12', CUT_NORMAL, f_in, fill_in, ctr, fmt='0.0')
put(wp, 'A13', '3 · Alta (≥)', f_b, fill_sec); put(wp, 'B13', CUT_ALTA, f_in, fill_in, ctr, fmt='0.0')
put(wp, 'A14', '4 · Urgente (≥)', f_b, fill_sec); put(wp, 'B14', CUT_URG, f_in, fill_in, ctr, fmt='0.0')
put(wp, 'C14', 'O corte proposto antes era 4,0. Com ele, Urgente ficava em 4,2% do total, abaixo da faixa mínima de 5%. Baixei para 3,9.', f_s, border=False, al=Alignment(wrap_text=False))
put(wp, 'A16', 'Regras que passam por cima da conta (categoria mínima)', f_b, border=False)
put(wp, 'A17', 'Risco = 5', f_b, fill_sec); put(wp, 'B17', 3, f_in, fill_in, ctr)
put(wp, 'A18', 'Para obra ou repasse da CEF = Sim', f_b, fill_sec); put(wp, 'B18', 4, f_in, fill_in, ctr)
put(wp, 'A20', 'Ajuste por natureza (somado à nota)', f_b, border=False)
NAT_ORDER = ['Controle', 'Documento normativo', 'Automação ou sistema', 'Definição ou decisão', 'Treinamento', 'Indicador']
for i, n in enumerate(NAT_ORDER):
    put(wp, 'A%d' % (21 + i), n, f_b, fill_sec); put(wp, 'B%d' % (21 + i), ADJ[n], f_in, fill_in, ctr, fmt='+0.0;-0.0;0.0')
put(wp, 'A28', 'Fator do processo', f_b, border=False)
put(wp, 'A29', 'Fator mínimo', f_b, fill_sec); put(wp, 'B29', 0.9, f_in, fill_in, ctr, fmt='0.00')
put(wp, 'A30', 'Fator máximo', f_b, fill_sec); put(wp, 'B30', 1.1, f_in, fill_in, ctr, fmt='0.00')
head(wp, 32, ['Processo', 'Dinheiro', 'Dependência', 'Irreversibilidade', 'Média', 'Fator'])
PROC_ORDER = list(PROC_NOTES.keys())
for i, p in enumerate(PROC_ORDER):
    r = 33 + i
    put(wp, 'A%d' % r, p, f_b, fill_sec)
    for col, v in zip('BCD', PROC_NOTES[p]): put(wp, '%s%d' % (col, r), v, f_in, fill_in, ctr)
    put(wp, 'E%d' % r, '=ROUND(AVERAGE(B%d:D%d),2)' % (r, r), f_n, fill_c, ctr, fmt='0.00')
    put(wp, 'F%d' % r, '=ROUND($B$29+($B$30-$B$29)*(E%d-1)/4,3)' % r, f_n, fill_c, ctr, fmt='0.000')
PR1, PR2 = 33, 32 + len(PROC_ORDER)

# =====================================================================
# TODAS
# =====================================================================
wt = wb.create_sheet('Todas')
wt.sheet_view.showGridLines = False
HR = 4
heads = ['Ordem\ngeral', 'ID', 'Atividade (link no ClickUp)', 'Processo', 'Responsável\nprincipal', 'Todos os\nresponsáveis', 'Status no\nClickUp', 'Situação', 'Natureza', 'Esforço\n(1–5)',
         'Risco de\nfraude e erro', 'Impacto\nfinanceiro', 'Obra e\nentrega', 'Legal, CEF e\ncompliance', 'Dependência', 'Para obra ou\nrepasse CEF?', 'Origem da\nnota',
         'Média\nponderada', 'Fator do\nprocesso', 'Ajuste da\nnatureza', 'Nota\nfinal', 'Cat. pela\nnota', 'Cat.\nfinal', 'Prioridade\nsugerida', 'Chave', 'Por que']
put(wt, 'A1', 'Todas as atividades do ClickUp com a prioridade sugerida', f_t, border=False, al=Alignment(vertical='center'))
put(wt, 'A2', 'Leitura do ClickUp em 27/10/2026 (todos os status). A prioridade que já estava no ClickUp não foi usada. Notas manuais e regra base são proposta minha; células amarelas podem ser editadas.', f_s, border=False)
head(wt, HR, heads, h=44)
widths = [7, 11, 62, 20, 16, 24, 15, 11, 20, 8, 10, 10, 8, 11, 11, 12, 14, 9, 9, 9, 8, 9, 7, 11, 9, 70]
for i, w in enumerate(widths): wt.column_dimensions[get_column_letter(i + 1)].width = w
wt.column_dimensions['Y'].hidden = True
FIRST = HR + 1; LAST = FIRST + len(rows) - 1
for i, o in enumerate(rows):
    r = FIRST + i; s = o['s']; t = o['t']
    put(wt, 'A%d' % r, '=IF(Y{r}="","",COUNTIF($Y${f}:$Y${l},">"&Y{r})+1)'.format(r=r, f=FIRST, l=LAST), f_b, fill_c, ctr)
    put(wt, 'B%d' % r, t['id'], f_g)
    c = put(wt, 'C%d' % r, re.sub(r'\s+', ' ', t['name']).strip(), f_l); c.hyperlink = t['url']
    put(wt, 'D%d' % r, o['proc']); put(wt, 'E%d' % r, o['resp1']); put(wt, 'F%d' % r, o['resps'])
    put(wt, 'G%d' % r, t['status']); put(wt, 'H%d' % r, o['sit'])
    put(wt, 'I%d' % r, s['nat'], f_in, fill_in)
    put(wt, 'J%d' % r, s['E'], f_in, fill_in, ctr)
    for col, k in zip('KLMNO', 'RFOLD'): put(wt, '%s%d' % (col, r), s[k], f_in, fill_in, ctr)
    put(wt, 'P%d' % r, s['cef'], f_in, fill_in, ctr)
    put(wt, 'Q%d' % r, o['origem'], f_g)
    put(wt, 'R%d' % r, '=ROUND((K{r}*Parametros!$B$4+L{r}*Parametros!$B$5+M{r}*Parametros!$B$6+N{r}*Parametros!$B$7+O{r}*Parametros!$B$8)/SUM(Parametros!$B$4:$B$8),2)'.format(r=r), f_n, fill_c, ctr, fmt='0.00')
    put(wt, 'S%d' % r, '=IFERROR(INDEX(Parametros!$F$%d:$F$%d,MATCH(D%d,Parametros!$A$%d:$A$%d,0)),1)' % (PR1, PR2, r, PR1, PR2), f_n, fill_c, ctr, fmt='0.000')
    put(wt, 'T%d' % r, '=IFERROR(INDEX(Parametros!$B$21:$B$26,MATCH(I%d,Parametros!$A$21:$A$26,0)),0)' % r, f_n, fill_c, ctr, fmt='+0.0;-0.0;0.0')
    put(wt, 'U%d' % r, '=ROUND(R{r}*S{r}+T{r},2)'.format(r=r), f_b, fill_c, ctr, fmt='0.00')
    put(wt, 'V%d' % r, '=IF(U{r}>=Parametros!$B$14,4,IF(U{r}>=Parametros!$B$13,3,IF(U{r}>=Parametros!$B$12,2,1)))'.format(r=r), f_n, fill_c, ctr)
    put(wt, 'W%d' % r, '=MAX(V{r},IF(K{r}=5,Parametros!$B$17,0),IF(P{r}="Sim",Parametros!$B$18,0))'.format(r=r), f_b, fill_c, ctr)
    put(wt, 'X%d' % r, '=CHOOSE(W{r},"Baixa","Normal","Alta","Urgente")'.format(r=r), f_b, fill_c, ctr)
    put(wt, 'Y%d' % r, '=W{r}*1000+U{r}*10-J{r}/10-ROW()/1000000'.format(r=r), f_g, fill_c, ctr, fmt='0.000000')
    put(wt, 'Z%d' % r, porque(o))
    wt.row_dimensions[r].height = 28
dv_n = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=False, showErrorMessage=True, errorTitle='Nota', error='Use um número inteiro de 1 a 5.')
dv_n.add('J%d:O%d' % (FIRST, LAST)); wt.add_data_validation(dv_n)
dv_s = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=False); dv_s.add('P%d:P%d' % (FIRST, LAST)); wt.add_data_validation(dv_s)
dv_g = DataValidation(type='list', formula1='=Parametros!$A$21:$A$26', allow_blank=False); dv_g.add('I%d:I%d' % (FIRST, LAST)); wt.add_data_validation(dv_g)
CL = {'Urgente': 'F4B6B6', 'Alta': 'FFE599', 'Normal': 'C9DAF8', 'Baixa': 'D9D9D9'}
for k, hexc in CL.items():
    wt.conditional_formatting.add('X%d:X%d' % (FIRST, LAST), FormulaRule(formula=['$X%d="%s"' % (FIRST, k)], fill=PatternFill('solid', start_color=hexc, end_color=hexc)))
wt.freeze_panes = 'D%d' % FIRST
wt.auto_filter.ref = 'A%d:Z%d' % (HR, LAST)

# =====================================================================
# URGENTE / ALTA (fotografia)
# =====================================================================
def lista(nome, cat, titulo, sub):
    ws = wb.create_sheet(nome); ws.sheet_view.showGridLines = False
    sel = [o for o in rows if o['cat'] == cat]
    sel.sort(key=lambda o: (0 if o['sit'] == 'Aberta' else 1, o['ordem']))
    n_ab = sum(1 for o in sel if o['sit'] == 'Aberta')
    put(ws, 'A1', titulo, f_t, border=False, al=Alignment(vertical='center'))
    put(ws, 'A2', '%s Total: %d (%d abertas e %d já concluídas ou canceladas). Fotografia do cálculo; para recalcular use a aba Todas.' % (sub, len(sel), n_ab, len(sel) - n_ab), f_s, border=False)
    hd = ['#', 'Ordem\ngeral', 'Atividade (link no ClickUp)', 'Processo', 'Responsável\nprincipal', 'Status no\nClickUp', 'Situação', 'Natureza', 'Nota\nfinal', 'Cat. pela\nnota', 'Veto\naplicado?',
          'Risco', 'Fin.', 'Obra', 'Legal', 'Dep.', 'Esforço', 'Por que entrou aqui']
    head(ws, 4, hd, h=36)
    for i, w in enumerate([5, 7, 62, 20, 16, 15, 11, 20, 8, 9, 10, 7, 6, 7, 7, 6, 8, 80]): ws.column_dimensions[get_column_letter(i + 1)].width = w
    for i, o in enumerate(sel, 1):
        r = 4 + i; s = o['s']; t = o['t']
        gray = o['sit'] != 'Aberta'
        fn = Font(name=F, size=10, color='7F7F7F') if gray else f_n
        fl = f_gl if gray else f_l
        put(ws, 'A%d' % r, i, fn, None, ctr); put(ws, 'B%d' % r, o['ordem'], fn, None, ctr)
        c = put(ws, 'C%d' % r, re.sub(r'\s+', ' ', t['name']).strip(), fl); c.hyperlink = t['url']
        put(ws, 'D%d' % r, o['proc'], fn); put(ws, 'E%d' % r, o['resp1'], fn); put(ws, 'F%d' % r, t['status'], fn); put(ws, 'G%d' % r, o['sit'], fn)
        put(ws, 'H%d' % r, s['nat'], fn); put(ws, 'I%d' % r, o['nota'], fn, None, ctr, fmt='0.00'); put(ws, 'J%d' % r, o['catn'], fn, None, ctr)
        put(ws, 'K%d' % r, 'Sim' if o['veto'] > o['catn'] else '', fn, None, ctr)
        for col, k in zip('LMNOP', 'RFOLD'): put(ws, '%s%d' % (col, r), s[k], fn, None, ctr)
        put(ws, 'Q%d' % r, s['E'], fn, None, ctr); put(ws, 'R%d' % r, porque(o), fn)
        ws.row_dimensions[r].height = 28
    ws.freeze_panes = 'D5'; ws.auto_filter.ref = 'A4:R%d' % (4 + len(sel))
    return sel
sel_u = lista('Urgente', 4, 'Sugestão: Urgente', 'Faixa sugerida de 5% a 10% do total.')
sel_a = lista('Alta', 3, 'Sugestão: Alta', 'Faixa sugerida de 15% a 25% do total.')

# =====================================================================
# RESUMO
# =====================================================================
wr = wb.create_sheet('Resumo'); wr.sheet_view.showGridLines = False
for k, v in {'A': 30, 'B': 13, 'C': 13, 'D': 13, 'E': 13, 'F': 22, 'G': 13, 'H': 13, 'I': 13, 'J': 13}.items(): wr.column_dimensions[k].width = v
rg = lambda col: 'Todas!$%s$%d:$%s$%d' % (col, FIRST, col, LAST)
put(wr, 'A1', 'Resumo: as faixas de prioridade', f_t, border=False, al=Alignment(vertical='center'))
put(wr, 'A2', 'Tudo calculado a partir da aba Todas. As faixas (azul) são a proposta de referência.', f_s, border=False)
put(wr, 'A4', 'Total de atividades', f_b, fill_sec); put(wr, 'B4', '=COUNTA(%s)' % rg('B'), f_b, fill_c, ctr)
put(wr, 'A5', 'Abertas hoje', f_b, fill_sec); put(wr, 'B5', '=COUNTIF(%s,"Aberta")' % rg('H'), f_b, fill_c, ctr)
put(wr, 'A6', 'Concluídas', f_b, fill_sec); put(wr, 'B6', '=COUNTIF(%s,"Concluída")' % rg('H'), f_b, fill_c, ctr)
put(wr, 'A7', 'Canceladas', f_b, fill_sec); put(wr, 'B7', '=COUNTIF(%s,"Cancelada")' % rg('H'), f_b, fill_c, ctr)
head(wr, 9, ['Categoria', 'Faixa mín.', 'Faixa máx.', 'Quantidade\n(total)', '% do total', 'Situação\nno total', 'Abertas', '% das\nabertas'], h=34)
cats = [('4 · Urgente', 4, 0.05, 0.10), ('3 · Alta', 3, 0.15, 0.25), ('2 · Normal', 2, 0.40, 0.50), ('1 · Baixa', 1, 0.20, 0.30)]
for i, (n, k, lo, hi) in enumerate(cats):
    r = 10 + i
    put(wr, 'A%d' % r, n, f_b, fill_sec)
    put(wr, 'B%d' % r, lo if k >= 3 else None, f_in, fill_in if k >= 3 else fill_c, ctr, fmt='0%')
    put(wr, 'C%d' % r, hi if k >= 3 else None, f_in, fill_in if k >= 3 else fill_c, ctr, fmt='0%')
    put(wr, 'D%d' % r, '=COUNTIF(%s,%d)' % (rg('W'), k), f_b, fill_c, ctr)
    put(wr, 'E%d' % r, '=D%d/$B$4' % r, f_n, fill_c, ctr, fmt='0.0%')
    if k >= 3: put(wr, 'F%d' % r, '=IF(E{r}<B{r},"abaixo da faixa",IF(E{r}>C{r},"acima da faixa","dentro da faixa"))'.format(r=r), f_n, fill_c, ctr)
    else: put(wr, 'F%d' % r, '—', f_n, fill_c, ctr)
    put(wr, 'G%d' % r, '=COUNTIFS(%s,%d,%s,"Aberta")' % (rg('W'), k, rg('H')), f_n, fill_c, ctr)
    put(wr, 'H%d' % r, '=G%d/$B$5' % r, f_n, fill_c, ctr, fmt='0.0%')
put(wr, 'A14', 'Urgente + Alta', f_b, fill_sec); put(wr, 'B14', None, f_n, fill_c); put(wr, 'C14', 0.30, f_in, fill_in, ctr, fmt='0%')
put(wr, 'D14', '=D10+D11', f_b, fill_c, ctr); put(wr, 'E14', '=D14/$B$4', f_n, fill_c, ctr, fmt='0.0%')
put(wr, 'F14', '=IF(E14>C14,"acima da trava","dentro da trava")', f_n, fill_c, ctr)
put(wr, 'G14', '=G10+G11', f_b, fill_c, ctr); put(wr, 'H14', '=G14/$B$5', f_n, fill_c, ctr, fmt='0.0%')
put(wr, 'A15', 'Trava de capacidade nas abertas (Urgente + Alta ≤)', f_b, fill_sec); put(wr, 'B15', None, f_n, fill_c); put(wr, 'C15', 0.30, f_in, fill_in, ctr, fmt='0%')
put(wr, 'D15', None, f_n, fill_c); put(wr, 'E15', None, f_n, fill_c)
put(wr, 'F15', '=IF(H14>C15,"acima da trava","dentro da trava")', f_b, fill_c, ctr)
put(wr, 'G15', None, f_n, fill_c); put(wr, 'H15', '=H14', f_n, fill_c, ctr, fmt='0.0%')
put(wr, 'A16', 'A faixa vale para o total (as concluídas mostram se a matriz concorda com o que a equipe já fez). A trava de capacidade vale para as abertas: é o que a equipe ainda precisa tocar.', f_s, border=False, al=Alignment(wrap_text=False))

head(wr, 18, ['Processo', 'Total', 'Abertas', 'Urgente\n(total)', 'Alta\n(total)', 'Urgente\n(abertas)', 'Alta\n(abertas)'], h=34)
for i, p in enumerate(PROC_ORDER):
    r = 19 + i
    put(wr, 'A%d' % r, p, f_b, fill_sec)
    put(wr, 'B%d' % r, '=COUNTIF(%s,A%d)' % (rg('D'), r), f_n, fill_c, ctr)
    put(wr, 'C%d' % r, '=COUNTIFS(%s,A%d,%s,"Aberta")' % (rg('D'), r, rg('H')), f_n, fill_c, ctr)
    put(wr, 'D%d' % r, '=COUNTIFS(%s,$A%d,%s,4)' % (rg('D'), r, rg('W')), f_n, fill_c, ctr)
    put(wr, 'E%d' % r, '=COUNTIFS(%s,$A%d,%s,3)' % (rg('D'), r, rg('W')), f_n, fill_c, ctr)
    put(wr, 'F%d' % r, '=COUNTIFS(%s,$A%d,%s,4,%s,"Aberta")' % (rg('D'), r, rg('W'), rg('H')), f_n, fill_c, ctr)
    put(wr, 'G%d' % r, '=COUNTIFS(%s,$A%d,%s,3,%s,"Aberta")' % (rg('D'), r, rg('W'), rg('H')), f_n, fill_c, ctr)
rr = 19 + len(PROC_ORDER)
put(wr, 'A%d' % rr, 'Total', f_b, fill_sec)
for col in 'BCDEFG': put(wr, '%s%d' % (col, rr), '=SUM(%s19:%s%d)' % (col, col, rr - 1), f_b, fill_c, ctr)

# por responsavel (abertas)
R0 = rr + 3
put(wr, 'A%d' % (R0 - 1), 'Urgentes abertas por responsável principal', f_b, border=False)
put(wr, 'D%d' % (R0 - 1), 'Limite por pessoa:', f_b, border=False, al=Alignment(horizontal='right'))
put(wr, 'E%d' % (R0 - 1), 3, f_in, fill_in, ctr)
head(wr, R0, ['Responsável', 'Urgente\n(abertas)', 'Alta\n(abertas)', 'Total\nabertas', 'Situação'], h=34)
cnt = collections.Counter(); cnta = collections.Counter(); tot = collections.Counter()
for o in rows:
    if o['sit'] == 'Aberta':
        tot[o['resp1']] += 1
        if o['cat'] == 4: cnt[o['resp1']] += 1
        if o['cat'] == 3: cnta[o['resp1']] += 1
resp_list = sorted(tot, key=lambda k: (-cnt[k], -cnta[k], k))
resp_list = [k for k in resp_list if cnt[k] + cnta[k] > 0][:25]
for i, nme in enumerate(resp_list):
    r = R0 + 1 + i
    put(wr, 'A%d' % r, nme, f_b, fill_sec)
    put(wr, 'B%d' % r, '=COUNTIFS(%s,A%d,%s,4,%s,"Aberta")' % (rg('E'), r, rg('W'), rg('H')), f_n, fill_c, ctr)
    put(wr, 'C%d' % r, '=COUNTIFS(%s,A%d,%s,3,%s,"Aberta")' % (rg('E'), r, rg('W'), rg('H')), f_n, fill_c, ctr)
    put(wr, 'D%d' % r, '=COUNTIFS(%s,A%d,%s,"Aberta")' % (rg('E'), r, rg('H')), f_n, fill_c, ctr)
    put(wr, 'E%d' % r, '=IF(B%d>$E$%d,"acima do limite","ok")' % (r, R0 - 1), f_n, fill_c, ctr)

# =====================================================================
# LEIA-ME
# =====================================================================
ws = wb['Sheet'] if 'Sheet' in wb.sheetnames else wb.create_sheet('Leia-me')
ws.title = 'Leia-me'; ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 30; ws.column_dimensions['C'].width = 110
put(ws, 'B2', 'Sugestão de prioridades: Urgente e Alta', f_t, border=False, al=Alignment(vertical='center'))
put(ws, 'B3', 'Projeto Next · Processos · lista "Implantação de Melhorias (5W2H)" · lida em 27/10/2026 · todas as atividades, de todos os status · proposta para validação', f_s, border=False)
n_all = len(rows); n_u = len(sel_u); n_a = len(sel_a)
n_man = sum(1 for o in rows if o['origem'].startswith('Manual'))
items = [
    ('O que é', 'A matriz de prioridades aplicada a todas as %d atividades do ClickUp (%d abertas, %d concluídas e %d canceladas). Mostra quais ficariam em Urgente e em Alta. A prioridade que já estava no ClickUp NÃO foi usada.' % (n_all, sum(1 for o in rows if o['sit'] == 'Aberta'), sum(1 for o in rows if o['sit'] == 'Concluída'), sum(1 for o in rows if o['sit'] == 'Cancelada'))),
    ('Resultado', 'Urgente: %d atividades (%.1f%% do total). Alta: %d atividades (%.1f%% do total). Faixas de referência: Urgente 5%% a 10%% e Alta 15%% a 25%%. Veja a aba Resumo.' % (n_u, 100 * n_u / n_all, n_a, 100 * n_a / n_all)),
    ('Como foi pontuado', 'Cada atividade recebe nota de 1 a 5 em cinco critérios: risco de fraude e erro (30%), impacto financeiro (20%), obra e entrega (20%), exigência legal/CEF/compliance (15%) e dependência (15%). A média é multiplicada pelo fator do processo (0,9 a 1,1) e somada ao ajuste da natureza (controle +0,3; indicador −0,3; treinamento −0,2). Esforço não entra na nota: só desempata.'),
    ('Cortes e vetos', 'Urgente a partir de 3,9; Alta a partir de 3,3; Normal a partir de 2,5. Risco 5 sobe a atividade para, no mínimo, Alta. "Para obra ou repasse da CEF" sobe para Urgente. O corte de Urgente era 4,0 na proposta anterior; com ele ficavam 41 atividades (4,2%), abaixo da faixa mínima de 5%. Baixei para 3,9.'),
    ('Quem pontuou', '%d atividades foram pontuadas por mim, uma a uma, a partir do nome, do processo, do status e do contexto que li do Projeto Next. As outras %d usaram uma regra automática por palavras do nome, limitada a nota 3 em cada critério, e só aparecem nas listas se isso bastar para chegar lá (a coluna "Origem da nota" mostra qual foi).' % (n_man, n_all - n_man)),
    ('Limites desta leitura', 'Li o NOME da atividade, não a descrição completa nem o histórico. As notas são PROPOSTA minha e devem ser contestadas pelo responsável e pelo gestor da área. Concluídas e canceladas entram para mostrar se a matriz concorda com o que a equipe fez; para decidir o que fazer agora, filtre Situação = Aberta. Atividades duplicadas entre Gadens e Make It recebem a mesma nota.'),
    ('Atenção', 'Entre as abertas, Urgente + Alta ficam acima da trava de capacidade de 30% (veja a aba Resumo). Isso indica excesso de coisa importante para a equipe tocar ao mesmo tempo, e não necessariamente erro da escala.'),
    ('Abas', 'Resumo: faixas, processos e responsáveis · Urgente e Alta: as listas sugeridas (abertas primeiro, concluídas e canceladas em cinza) · Todas: as atividades com a conta (fórmulas vivas) · Parametros: pesos, cortes, vetos e fatores editáveis.'),
    ('Cores', 'Azul em fundo amarelo = pode editar. Fundo cinza = fórmula. Prioridade com as cores das bandeiras do ClickUp: vermelho Urgente, amarelo Alta, azul Normal, cinza Baixa.'),
]
r = 5
for k, v in items:
    put(ws, 'B%d' % r, k, f_b, fill_sec); put(ws, 'C%d' % r, v)
    ws.row_dimensions[r].height = max(30, 15 * (len(v) // 110 + 1 + v.count('\n')))
    r += 1

# ordem das abas
order = ['Leia-me', 'Resumo', 'Urgente', 'Alta', 'Todas', 'Parametros']
wb._sheets = [wb[n] for n in order]
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('salvo', OUT, '| Urgente', n_u, 'Alta', n_a, '| manuais', n_man, 'de', n_all)
json_out = {o['t']['id']: (o['nota'], o['cat'], o['ordem']) for o in rows}
import json; json.dump(json_out, open('esperado.json', 'w'))
