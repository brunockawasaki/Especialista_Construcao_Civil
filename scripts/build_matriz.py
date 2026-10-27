# -*- coding: utf-8 -*-
import json, re, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

SRC = os.path.join(os.path.dirname(__file__), 'clickup_abertas.json')
OUT = sys.argv[1]

tasks = [t for t in json.load(open(SRC, encoding='utf-8')) if t['status'] not in ('cancelado', 'concluído')]

# ---------- estilos ----------
F = 'Arial'
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True, color='1F3864')
f_sub = Font(name=F, size=10, italic=True, color='595959')
f_head = Font(name=F, size=10, bold=True, color='FFFFFF')
f_input = Font(name=F, size=10, color='0000FF')
f_link = Font(name=F, size=10, color='0563C1', underline='single')
f_grey = Font(name=F, size=9, color='7F7F7F')
fill_head = PatternFill('solid', start_color='1F3864')
fill_in = PatternFill('solid', start_color='FFF2CC')
fill_sec = PatternFill('solid', start_color='D9E1F2')
fill_calc = PatternFill('solid', start_color='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def put(ws, ref, v, font=f_norm, fill=None, al=wrap, border=True, fmt=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fill: c.fill = fill
    c.alignment = al
    if border: c.border = box
    if fmt: c.number_format = fmt
    return c

def head(ws, row, labels, col0=1):
    for i, l in enumerate(labels):
        put(ws, '%s%d' % (get_column_letter(col0 + i), row), l, f_head, fill_head, center)

wb = Workbook()

# =====================================================================
# LEIA-ME
# =====================================================================
ws = wb.active
ws.title = 'Leia-me'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 105
put(ws, 'B2', 'Matriz de Prioridades das Atividades do ClickUp', f_title, border=False, al=Alignment(vertical='center'))
put(ws, 'B3', 'Projeto Next · Processos · lista "Implantação de Melhorias (5W2H)" · proposta para validação', f_sub, border=False)
rows = [
    ('Para que serve', 'Dar uma nota de importância a cada atividade aberta do ClickUp e traduzir a nota em 4 categorias: 1 Baixa, 2 Normal, 3 Alta, 4 Urgente. A prioridade que hoje está no ClickUp NÃO foi usada.'),
    ('A ideia em uma frase', 'Importância é uma coisa, prazo é outra e esforço é outra. A nota mede só a importância. O esforço e o prazo servem para ordenar as atividades dentro de cada categoria.'),
    ('Como a nota é calculada', '1) Cada atividade recebe nota de 1 a 5 em 5 critérios (aba Matriz).\n2) A média é ponderada pelos pesos da aba Matriz.\n3) Multiplica-se pelo fator do processo (aba Processos): um ajuste pequeno, de 0,9 a 1,1.\n4) Soma-se o ajuste da natureza da ação (controle sobe, indicador desce).\n5) A nota final vira categoria pelos cortes da aba Matriz.\n6) Duas regras passam por cima da conta (vetos): risco de fraude nota 5 vira, no mínimo, Alta; atividade que para a obra ou o repasse da CEF vira Urgente.'),
    ('Como usar a aba Tarefas', '1) Preencha as células amarelas: esforço (H), as 5 notas (I a M) e, se for o caso, "Para obra ou repasse CEF?" (N).\n2) As colunas cinza calculam sozinhas: nota, categoria e prioridade.\n3) A coluna "Ordem geral" diz o que fazer primeiro: categoria, depois nota, depois menor esforço.\n4) Dá para filtrar por processo, responsável, status ou prioridade.'),
    ('Cores', 'Texto azul em fundo amarelo = pode editar (pesos, cortes, notas, esforço).\nFundo cinza = fórmula, não editar.\nFundo azul-escuro = títulos.'),
    ('O que é proposta minha', 'Todos os pesos, cortes, ajustes e médias de processo são PROPOSTA para validar. Nenhum veio de dado oficial. As notas das 12 atividades marcadas como "Âncora (proposta)" também são minhas: servem de exemplo e de referência para calibrar. Troque qualquer valor azul e tudo recalcula.'),
    ('Fonte dos dados', 'ClickUp, lista "Implantação de Melhorias (5W2H)", lida em 21/10/2026. Entram %d atividades abertas (status: sem prazo definido, em progresso, em atraso e pausado). Ficam de fora as concluídas e as canceladas. Processo e Tipo de Ação vêm dos campos do ClickUp; onde o campo estava vazio usei a etiqueta; sem nenhum dos dois ficou "Sem processo definido".' % len(tasks)),
    ('Sugestões automáticas', 'A coluna "Natureza" (Controle, Documento normativo, etc.) é uma sugestão minha a partir do tipo de ação do ClickUp e de palavras do nome da tarefa. Revise e troque pela lista suspensa se estiver errada.'),
    ('Limites', 'A planilha não altera nada no ClickUp. As notas dos critérios (colunas I a M) ficam em branco para as áreas e o Processos pontuarem, exceto nas âncoras.'),
]
r = 5
for k, v in rows:
    put(ws, 'B%d' % r, k, f_bold, fill_sec)
    put(ws, 'C%d' % r, v)
    ws.row_dimensions[r].height = max(30, 15 * (v.count('\n') + 1 + len(v) // 120))
    r += 1

# =====================================================================
# MATRIZ
# =====================================================================
wm = wb.create_sheet('Matriz')
wm.sheet_view.showGridLines = False
widths = {'A': 30, 'B': 44, 'C': 11, 'D': 30, 'E': 30, 'F': 30, 'G': 30, 'H': 30}
for k, v in widths.items(): wm.column_dimensions[k].width = v
put(wm, 'A1', 'Matriz de critérios, pesos e regras', f_title, border=False, al=Alignment(vertical='center'))
put(wm, 'A2', 'Células azuis em fundo amarelo podem ser editadas. Pesos e cortes são proposta.', f_sub, border=False)
head(wm, 4, ['Critério', 'Pergunta', 'Peso (%)', 'Nota 1', 'Nota 2', 'Nota 3', 'Nota 4', 'Nota 5'])
crit = [
    ('Risco de fraude e erro', 'Pode haver desvio de dinheiro ou erro grave sem que alguém perceba?', 30,
     'Nenhum risco relevante.', 'Erro pequeno e fácil de achar.', 'Erro ou desvio possível, mas há conferência depois.',
     'Desvio ou erro de valor possível e sem barreira clara.', 'Dinheiro sai ou dado crítico muda sem trava, aprovação ou conferência.'),
    ('Impacto financeiro', 'Quanto dinheiro está em jogo ou pode ser perdido ou ganho?', 20,
     'Sem efeito em caixa.', 'Efeito pequeno (exemplo: até R$ 10 mil).', 'Efeito médio (exemplo: R$ 10 a 50 mil).',
     'Efeito grande (exemplo: R$ 50 a 200 mil).', 'Efeito muito grande (exemplo: acima de R$ 200 mil). Faixas a confirmar com o Financeiro.'),
    ('Impacto na obra e na entrega', 'Atrasa obra, lançamento, repasse ou entrega?', 20,
     'Nenhum efeito.', 'Atrasa rotina interna.', 'Atrasa uma etapa, sem afetar marco.',
     'Pode atrasar marco de obra ou de lançamento.', 'Para a obra, o repasse da CEF ou a entrega.'),
    ('Exigência legal, CEF e compliance', 'Existe lei, norma, contrato, CEF ou LGPD exigindo isso?', 15,
     'Nenhuma.', 'Boa prática interna.', 'Exigência interna formal (política, procedimento).',
     'Exigência de contrato, CEF ou LGPD.', 'Lei, norma de segurança ou risco de multa ou embargo.'),
    ('Dependência', 'Outras atividades ou áreas esperam por esta?', 15,
     'Ninguém espera.', 'Uma atividade espera.', 'Duas ou três atividades, ou uma área.',
     'Várias áreas esperam.', 'Destrava um processo inteiro ou vários fluxos.'),
]
for i, c in enumerate(crit):
    rr = 5 + i
    put(wm, 'A%d' % rr, c[0], f_bold, fill_sec)
    put(wm, 'B%d' % rr, c[1])
    put(wm, 'C%d' % rr, c[2], f_input, fill_in, center)
    for j in range(5):
        put(wm, '%s%d' % ('DEFGH'[j], rr), c[3 + j])
    wm.row_dimensions[rr].height = 58
put(wm, 'A10', 'Soma dos pesos', f_bold, fill_calc)
put(wm, 'B10', '=IF(C10=100,"Soma correta (100%)","ATENÇÃO: os pesos devem somar 100")', f_norm, fill_calc)
put(wm, 'C10', '=SUM(C5:C9)', f_bold, fill_calc, center)
put(wm, 'A12', 'Esforço (não entra na nota)', f_bold, fill_sec)
put(wm, 'B12', 'Quanto custa fazer? Só ordena dentro da categoria e marca ganho rápido.', f_norm)
put(wm, 'C12', '—', f_norm, fill_calc, center)
for j, t in enumerate(['Até 1 semana.', 'Até 1 mês.', 'Até 3 meses.', 'Mais de 3 meses.', 'Depende de contratação, sistema externo ou decisão da diretoria.']):
    put(wm, '%s12' % 'DEFGH'[j], 'Nível %d: %s' % (j + 1, t))
wm.row_dimensions[12].height = 45

put(wm, 'A14', 'Cortes das categorias', f_bold, border=False)
head(wm, 15, ['Categoria', 'Significado', 'Nota mínima (≥)'])
cuts = [('1 · Baixa', 'Pode esperar. Fazer quando houver folga.', 0), ('2 · Normal', 'Entra na rotina de melhoria.', 2.5),
        ('3 · Alta', 'Planejar para o próximo ciclo.', 3.3), ('4 · Urgente', 'Tratar já, com dono e prazo.', 4.0)]
for i, (a, b, c) in enumerate(cuts):
    rr = 16 + i
    put(wm, 'A%d' % rr, a, f_bold, fill_sec); put(wm, 'B%d' % rr, b)
    if i == 0:
        put(wm, 'C%d' % rr, c, f_norm, fill_calc, center, fmt='0.0')
    else:
        put(wm, 'C%d' % rr, c, f_input, fill_in, center, fmt='0.0')

put(wm, 'A21', 'Regras que passam por cima da conta (vetos)', f_bold, border=False)
head(wm, 22, ['Regra', 'Quando vale', 'Categoria mínima'])
put(wm, 'A23', 'Risco no nível máximo', f_bold, fill_sec); put(wm, 'B23', 'A nota de "Risco de fraude e erro" é 5.'); put(wm, 'C23', 3, f_input, fill_in, center)
put(wm, 'A24', 'Para obra ou repasse da CEF', f_bold, fill_sec); put(wm, 'B24', 'Coluna N da aba Tarefas marcada como "Sim".'); put(wm, 'C24', 4, f_input, fill_in, center)

put(wm, 'A26', 'Ajuste por natureza da ação (somado à nota)', f_bold, border=False)
head(wm, 27, ['Natureza', 'Quando usar', 'Ajuste'])
nat = [('Controle', 'Alçada, trava, segregação, aprovação, conciliação, acesso. Protege contra fraude e erro.', 0.3),
       ('Documento normativo', 'POP, IT, Política ou Regimento. A nota vem do risco que o documento cobre.', 0.0),
       ('Automação ou sistema', 'Integração, parametrização, campo, módulo. A nota vem do que protege ou acelera.', 0.0),
       ('Definição ou decisão', 'Definir regra, responsável ou prazo, ou avaliar uma opção.', 0.0),
       ('Treinamento', 'Treinar equipe. Sobe a nota pelos critérios quando sustenta um controle.', -0.2),
       ('Indicador', 'Medir sem proteger. Se o indicador detecta desvio (exemplo: avanço físico x financeiro), classificar como Controle.', -0.3)]
for i, (a, b, c) in enumerate(nat):
    rr = 28 + i
    put(wm, 'A%d' % rr, a, f_bold, fill_sec); put(wm, 'B%d' % rr, b); put(wm, 'C%d' % rr, c, f_input, fill_in, center, fmt='+0.0;-0.0;0.0')
    wm.row_dimensions[rr].height = 40
put(wm, 'A35', 'Nota final = média ponderada dos 5 critérios × fator do processo + ajuste da natureza. Depois aplicam-se os vetos. Esforço e prazo só ordenam.', f_sub, border=False)

# =====================================================================
# PROCESSOS
# =====================================================================
wp = wb.create_sheet('Processos')
wp.sheet_view.showGridLines = False
for k, v in {'A': 28, 'B': 13, 'C': 13, 'D': 16, 'E': 10, 'F': 10, 'G': 70}.items(): wp.column_dimensions[k].width = v
put(wp, 'A1', 'Peso dos processos (cadeia de valor)', f_title, border=False, al=Alignment(vertical='center'))
put(wp, 'A2', 'Três perguntas, de 1 a 5, por processo. O fator é um ajuste pequeno sobre a nota; o risco da atividade manda mais do que o processo.', f_sub, border=False)
put(wp, 'A4', 'Fator mínimo', f_bold, fill_sec); put(wp, 'B4', 0.9, f_input, fill_in, center, fmt='0.00')
put(wp, 'A5', 'Fator máximo', f_bold, fill_sec); put(wp, 'B5', 1.1, f_input, fill_in, center, fmt='0.00')
head(wp, 7, ['Processo', 'Dinheiro (1–5)', 'Dependência (1–5)', 'Irreversibilidade (1–5)', 'Média', 'Fator', 'Por quê (hipótese para validar)'])
procs = [
    ('Incorporação Make It', 5, 5, 5, 'Define produto, projeto e custo. Quase tudo depende e o erro volta caro.'),
    ('Incorporação Gadens', 5, 5, 5, 'Mesma lógica; ciclo mais longo e valor por unidade maior.'),
    ('Engenharia', 5, 5, 4, 'Orçamento, planejamento e obra concentram o maior custo e o caminho crítico.'),
    ('Novos Negócios', 5, 3, 5, 'Poucas decisões, mas irreversíveis (compra de terreno, contratos, SPE).'),
    ('Suprimentos', 5, 4, 3, 'Maior parte do custo de obra passa por compras; erro de compra tem correção parcial.'),
    ('Financeiro', 4, 5, 3, 'Todo pagamento e conciliação passa por aqui. É processo de apoio, mas de alto risco.'),
    ('GMO', 4, 4, 3, 'Contratos de empreiteiros e medição; a obra depende deles para iniciar e pagar.'),
    ('Comercial Make It', 4, 4, 3, 'Caixa depende do ritmo de venda e do repasse da CEF.'),
    ('Comercial Gadens', 4, 3, 3, 'Ciclo mais lento e com mais folga que o Make It.'),
    ('Sem processo definido', 3, 3, 3, 'Atividade sem Processo nem etiqueta no ClickUp. Fator neutro.'),
]
for i, (n, a, b, c, why) in enumerate(procs):
    rr = 8 + i
    put(wp, 'A%d' % rr, n, f_bold, fill_sec)
    for col, v in zip('BCD', (a, b, c)):
        put(wp, '%s%d' % (col, rr), v, f_input, fill_in, center)
    put(wp, 'E%d' % rr, '=ROUND(AVERAGE(B%d:D%d),2)' % (rr, rr), f_norm, fill_calc, center, fmt='0.00')
    put(wp, 'F%d' % rr, '=ROUND($B$4+($B$5-$B$4)*(E%d-1)/4,3)' % rr, f_norm, fill_calc, center, fmt='0.000')
    put(wp, 'G%d' % rr, why)
    wp.row_dimensions[rr].height = 30
PROC_FIRST, PROC_LAST = 8, 7 + len(procs)

# =====================================================================
# TAREFAS
# =====================================================================
wt = wb.create_sheet('Tarefas')
wt.sheet_view.showGridLines = False

def norm_proc(t):
    p = re.sub(r'\s+', ' ', (t['custom'].get('Processo') or '')).strip()
    known = {x[0] for x in procs}
    if p in known: return p
    if p.startswith('Engenharia'): return 'Engenharia'
    tagmap = {'novos negócios': 'Novos Negócios', 'incorporação make it': 'Incorporação Make It', 'incorporação gadens': 'Incorporação Gadens',
              'engenharia': 'Engenharia', 'gmo': 'GMO', 'suprimentos': 'Suprimentos', 'financeiro': 'Financeiro',
              'comercial gadens': 'Comercial Gadens', 'comercial make it': 'Comercial Make It'}
    for g in t['tags']:
        if g in tagmap: return tagmap[g]
    return 'Sem processo definido'

def natureza(t):
    tipo = re.sub(r'\s+', ' ', (t['custom'].get('Tipo de Ação') or '')).strip().lower()
    n = t['name'].lower()
    if tipo == 'documento normativo' or re.match(r'^\W*elaborar (pop|it|política|regimento)', n): return 'Documento normativo'
    if tipo == 'indicador' or re.match(r'^\W*implementar indicador|^\W*índice de', n): return 'Indicador'
    if 'treinamento' in n or 'treinar' in n: return 'Treinamento'
    if re.search(r'alçada|trava|bloque|segrega|perfis? de acesso|concilia|due diligence|aprovação de|dados banc|homologa|auditoria|catraca|acesso', n): return 'Controle'
    if re.search(r'automat|integra|parametr|criar campo|habilitar|módulo|rpa|\bbi\b|clicksign|approvo|\bcv\b|mega|prevision|robô|plataforma|sistema|g-control', n): return 'Automação ou sistema'
    return 'Definição ou decisão'

anchors = {  # id: (esforco, risco, fin, obra, legal, dep, cef, natureza, por que)
    '86acmweg8': (4, 5, 4, 3, 3, 5, 'Não', 'Controle', 'Sem alçada formal, vários fluxos decidem "conforme alçada" sem regra. Difícil de fazer (depende da diretoria).'),
    '86ahckp0b': (3, 4, 3, 5, 5, 3, 'Não', None, 'Procedimento para parar obra por segurança: risco à vida, norma de segurança e efeito direto na obra.'),
    '86aj81736': (3, 4, 4, 4, 2, 2, 'Não', 'Controle', 'Compara o avanço físico com o financeiro. Detecta pagamento acima do executado, por isso conta como Controle.'),
    '86ahket8w': (2, 5, 4, 1, 2, 2, 'Não', 'Controle', 'Trava de dados bancários obrigatórios: sem ela, o pagamento pode sair para conta errada. Esforço baixo.'),
    '86aghkwdz': (3, 5, 3, 1, 3, 3, 'Não', 'Controle', 'Revisão de perfis de acesso: sem ela, o mesmo usuário pode pedir e aprovar (segregação de funções).'),
    '86ahxbzgc': (3, 5, 4, 1, 2, 2, 'Não', None, 'POP de pagamento manual: caminho em que o dinheiro sai fora da remessa automática.'),
    '17tpepcgphr': (4, 4, 4, 1, 2, 3, 'Não', 'Controle', 'Controle da Política Comercial: sem ele, condições de venda podem sair da regra.'),
    '17tpepcgkv5': (3, 2, 3, 1, 3, 1, 'Não', None, 'Avaliar uma modalidade nova de venda: decisão, sem risco imediato.'),
    '86aghm6gm': (1, 1, 1, 1, 1, 1, 'Não', None, 'Avaliar Uber corporativo: efeito mínimo.'),
    '86aghg2zx': (2, 5, 3, 1, 3, 2, 'Não', 'Controle', 'Regra de conta PJ para serviço PJ: evita pagar prestador em conta de terceiro.'),
    '86aj7rq96': (3, 3, 3, 3, 1, 2, 'Não', None, 'Política de almoxarifado e inventário: perda de material e acuracidade de estoque.'),
    '17tpepcpr24': (2, 3, 3, 1, 2, 4, 'Não', None, 'Campo de permutante e tipos de permuta no CV: vários fluxos de permuta dependem dele.'),
}

HEAD_ROW = 4
heads = ['ID', 'Tarefa', 'Processo', 'Responsável', 'Status', 'Prazo', 'Natureza', 'Esforço\n(1–5)',
         'Risco de\nfraude e erro', 'Impacto\nfinanceiro', 'Obra e\nentrega', 'Legal, CEF\ne compliance', 'Dependência',
         'Para obra ou\nrepasse CEF?', 'Média\nponderada', 'Fator do\nprocesso', 'Ajuste da\nnatureza', 'Nota\nfinal',
         'Categoria\npela nota', 'Categoria\nfinal', 'Prioridade', 'Chave', 'Ordem\ngeral', 'Origem da nota', 'Por que (âncora)']
put(wt, 'A1', 'Atividades abertas do ClickUp com a matriz', f_title, border=False, al=Alignment(vertical='center'))
put(wt, 'A2', 'Preencha as células amarelas (H a N). Colunas cinza calculam sozinhas. Âncoras: 12 atividades já pontuadas como proposta, para calibrar.', f_sub, border=False)
head(wt, HEAD_ROW, heads)
wt.row_dimensions[HEAD_ROW].height = 42
colw = [11, 60, 20, 22, 14, 11, 20, 8, 11, 10, 9, 12, 12, 13, 10, 9, 9, 8, 10, 9, 11, 8, 8, 20, 55]
for i, w in enumerate(colw): wt.column_dimensions[get_column_letter(i + 1)].width = w
wt.column_dimensions['V'].hidden = True

order_status = {'em atraso': 0, 'em progresso': 1, 'sem prazo definido': 2, 'pausado': 3}
tasks.sort(key=lambda t: (norm_proc(t), order_status.get(t['status'], 9), t['name']))
FIRST = HEAD_ROW + 1
LAST = FIRST + len(tasks) - 1
anchors_found = []
for i, t in enumerate(tasks):
    r = FIRST + i
    a = anchors.get(t['id'])
    put(wt, 'A%d' % r, t['id'], f_grey)
    c = put(wt, 'B%d' % r, t['name'], f_link)
    c.hyperlink = t['url']
    put(wt, 'C%d' % r, norm_proc(t))
    put(wt, 'D%d' % r, ', '.join(t['assignees']) or '(sem responsável)')
    put(wt, 'E%d' % r, t['status'])
    due = None
    if t.get('due'):
        due = (datetime.datetime.utcfromtimestamp(int(t['due']) / 1000) - datetime.timedelta(hours=3)).date()
    put(wt, 'F%d' % r, due, f_norm, None, center, fmt='dd/mm/yyyy')
    nt = (a[7] if a and a[7] else natureza(t))
    put(wt, 'G%d' % r, nt, f_input, fill_in)
    vals = [None] * 5 + ['']
    esf = None
    if a:
        anchors_found.append(t['id'])
        esf = a[0]; vals = list(a[1:6]) + [a[6]]
    put(wt, 'H%d' % r, esf, f_input, fill_in, center)
    for col, v in zip('IJKLM', vals[:5]):
        put(wt, '%s%d' % (col, r), v, f_input, fill_in, center)
    put(wt, 'N%d' % r, vals[5] if a else None, f_input, fill_in, center)
    put(wt, 'O%d' % r, '=IF(COUNT(I{r}:M{r})<5,"",ROUND((I{r}*Matriz!$C$5+J{r}*Matriz!$C$6+K{r}*Matriz!$C$7+L{r}*Matriz!$C$8+M{r}*Matriz!$C$9)/SUM(Matriz!$C$5:$C$9),2))'.format(r=r), f_norm, fill_calc, center, fmt='0.00')
    put(wt, 'P%d' % r, '=IFERROR(INDEX(Processos!$F$%d:$F$%d,MATCH(C%d,Processos!$A$%d:$A$%d,0)),1)' % (PROC_FIRST, PROC_LAST, r, PROC_FIRST, PROC_LAST), f_norm, fill_calc, center, fmt='0.000')
    put(wt, 'Q%d' % r, '=IFERROR(INDEX(Matriz!$C$28:$C$33,MATCH(G%d,Matriz!$A$28:$A$33,0)),0)' % r, f_norm, fill_calc, center, fmt='+0.0;-0.0;0.0')
    put(wt, 'R%d' % r, '=IF(O{r}="","",ROUND(O{r}*P{r}+Q{r},2))'.format(r=r), f_bold, fill_calc, center, fmt='0.00')
    put(wt, 'S%d' % r, '=IF(R{r}="","",IF(R{r}>=Matriz!$C$19,4,IF(R{r}>=Matriz!$C$18,3,IF(R{r}>=Matriz!$C$17,2,1))))'.format(r=r), f_norm, fill_calc, center)
    put(wt, 'T%d' % r, '=IF(S{r}="","",MAX(S{r},IF(I{r}=5,Matriz!$C$23,0),IF(N{r}="Sim",Matriz!$C$24,0)))'.format(r=r), f_bold, fill_calc, center)
    put(wt, 'U%d' % r, '=IF(T{r}="","",CHOOSE(T{r},"Baixa","Normal","Alta","Urgente"))'.format(r=r), f_bold, fill_calc, center)
    put(wt, 'V%d' % r, '=IF(T{r}="","",T{r}*1000+R{r}*10-IF(H{r}="",3,H{r})/10-ROW()/1000000)'.format(r=r), f_grey, fill_calc, center, fmt='0.000000')
    put(wt, 'W%d' % r, '=IF(V{r}="","",COUNTIF($V${f}:$V${l},">"&V{r})+1)'.format(r=r, f=FIRST, l=LAST), f_norm, fill_calc, center)
    put(wt, 'X%d' % r, 'Âncora (proposta)' if a else None, f_norm)
    put(wt, 'Y%d' % r, a[8] if a else None, f_norm)
    if a: wt.row_dimensions[r].height = 40
    else: wt.row_dimensions[r].height = 28

# validacoes
dv_n = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True, showErrorMessage=True,
                      errorTitle='Nota', error='Use um número inteiro de 1 a 5.')
dv_n.add('H%d:M%d' % (FIRST, LAST)); wt.add_data_validation(dv_n)
dv_s = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); dv_s.add('N%d:N%d' % (FIRST, LAST)); wt.add_data_validation(dv_s)
dv_g = DataValidation(type='list', formula1='=Matriz!$A$28:$A$33', allow_blank=True); dv_g.add('G%d:G%d' % (FIRST, LAST)); wt.add_data_validation(dv_g)

# cores da prioridade
red = PatternFill('solid', start_color='F4B6B6', end_color='F4B6B6')
org = PatternFill('solid', start_color='F9CB9C', end_color='F9CB9C')
yel = PatternFill('solid', start_color='FFE599', end_color='FFE599')
grn = PatternFill('solid', start_color='B6D7A8', end_color='B6D7A8')
rng = 'U%d:U%d' % (FIRST, LAST)
wt.conditional_formatting.add(rng, FormulaRule(formula=['$U%d="Urgente"' % FIRST], fill=red))
wt.conditional_formatting.add(rng, FormulaRule(formula=['$U%d="Alta"' % FIRST], fill=org))
wt.conditional_formatting.add(rng, FormulaRule(formula=['$U%d="Normal"' % FIRST], fill=yel))
wt.conditional_formatting.add(rng, FormulaRule(formula=['$U%d="Baixa"' % FIRST], fill=grn))
wt.freeze_panes = 'C%d' % FIRST
wt.auto_filter.ref = 'A%d:Y%d' % (HEAD_ROW, LAST)

# =====================================================================
# RESUMO
# =====================================================================
wr = wb.create_sheet('Resumo')
wr.sheet_view.showGridLines = False
for k, v in {'A': 30, 'B': 14, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 16}.items(): wr.column_dimensions[k].width = v
put(wr, 'A1', 'Resumo da priorização', f_title, border=False, al=Alignment(vertical='center'))
put(wr, 'A2', 'Tudo calculado a partir da aba Tarefas.', f_sub, border=False)
rg = lambda col: 'Tarefas!$%s$%d:$%s$%d' % (col, FIRST, col, LAST)
put(wr, 'A4', 'Atividades abertas', f_bold, fill_sec); put(wr, 'B4', '=COUNTA(%s)' % rg('A'), f_bold, fill_calc, center)
put(wr, 'A5', 'Já pontuadas', f_bold, fill_sec); put(wr, 'B5', '=COUNT(%s)' % rg('R'), f_bold, fill_calc, center)
put(wr, 'A6', 'Falta pontuar', f_bold, fill_sec); put(wr, 'B6', '=B4-B5', f_bold, fill_calc, center)
head(wr, 8, ['Prioridade', 'Quantidade', '% das pontuadas', 'Meta máxima', 'Situação'])
meta = [('4 · Urgente', 4, 0.10), ('3 · Alta', 3, 0.25), ('2 · Normal', 2, None), ('1 · Baixa', 1, None)]
for i, (n, k, m) in enumerate(meta):
    rr = 9 + i
    put(wr, 'A%d' % rr, n, f_bold, fill_sec)
    put(wr, 'B%d' % rr, '=COUNTIF(%s,%d)' % (rg('T'), k), f_norm, fill_calc, center)
    put(wr, 'C%d' % rr, '=IF($B$5=0,0,B%d/$B$5)' % rr, f_norm, fill_calc, center, fmt='0%')
    if m is not None:
        put(wr, 'D%d' % rr, m, f_input, fill_in, center, fmt='0%')
        put(wr, 'E%d' % rr, '=IF($B$5=0,"—",IF(C%d>D%d,"Acima da meta: revisar cortes","Dentro da meta"))' % (rr, rr), f_norm, fill_calc, center)
    else:
        put(wr, 'D%d' % rr, '—', f_norm, fill_calc, center); put(wr, 'E%d' % rr, '—', f_norm, fill_calc, center)
put(wr, 'A13', 'Meta máxima: se Urgente passar de 10% ou Alta de 25% das pontuadas, a escala perde sentido. É proposta; edite as células azuis.', f_sub, border=False)
head(wr, 15, ['Processo', 'Abertas', 'Pontuadas', 'Urgente', 'Alta', 'Normal', 'Baixa'])
for i, (n, *_r) in enumerate(procs):
    rr = 16 + i
    put(wr, 'A%d' % rr, n, f_bold, fill_sec)
    put(wr, 'B%d' % rr, '=COUNTIF(%s,A%d)' % (rg('C'), rr), f_norm, fill_calc, center)
    put(wr, 'C%d' % rr, '=COUNTIFS(%s,A%d,%s,">0")' % (rg('C'), rr, rg('R')), f_norm, fill_calc, center)
    for col, k in zip('DEFG', (4, 3, 2, 1)):
        put(wr, '%s%d' % (col, rr), '=COUNTIFS(%s,$A%d,%s,%d)' % (rg('C'), rr, rg('T'), k), f_norm, fill_calc, center)
rr = 16 + len(procs)
put(wr, 'A%d' % rr, 'Total', f_bold, fill_sec)
for col in 'BCDEFG':
    put(wr, '%s%d' % (col, rr), '=SUM(%s16:%s%d)' % (col, col, rr - 1), f_bold, fill_calc, center)

wb.move_sheet('Resumo', offset=-1)  # Resumo antes de Tarefas
wb.save(OUT)
print('salvo', OUT, '| tarefas', len(tasks), '| ancoras encontradas', len(anchors_found), 'de', len(anchors))
print('ancoras ausentes:', [k for k in anchors if k not in anchors_found])
